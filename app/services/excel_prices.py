from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.repo import Repo

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)


def _settlement_header(raw) -> str:
    if raw is None:
        return ""
    text = str(raw).strip()
    if not text or text.lower() == "транспорт / поселок":
        return ""
    if " / " in text:
        text = text.split(" / ", 1)[0].strip()
    return text


def export_prices_xlsx(repo: Repo, output_path: Path) -> Path:
    vehicles = repo.list_vehicles()
    settlements = repo.list_settlements()
    price_map = {(p.vehicle_id, p.settlement_id): p.amount for p in repo.list_prices()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"

    # A1 — угол; сверху поселки, слева транспорт
    ws.cell(1, 1, "Транспорт / Поселок")
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).font = HEADER_FONT

    for col_idx, settlement in enumerate(settlements, start=2):
        cell = ws.cell(1, col_idx, settlement.name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row_idx, vehicle in enumerate(vehicles, start=2):
        cell = ws.cell(row_idx, 1, vehicle.name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        for col_idx, settlement in enumerate(settlements, start=2):
            amount = price_map.get((vehicle.id, settlement.id))
            ws.cell(row_idx, col_idx, amount if amount is not None else "")

    ws.column_dimensions["A"].width = 28
    for col_idx in range(2, len(settlements) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def import_prices_xlsx(repo: Repo, source: Path | bytes) -> dict:
    """
    Импорт матрицы: новые строки → транспорт, новые столбцы → поселки.
    Заголовок столбца — имя поселка. Старый вид «Поселок / Участок» берёт только поселок.
    """
    if isinstance(source, bytes):
        wb = load_workbook(BytesIO(source), data_only=True)
    else:
        wb = load_workbook(source, data_only=True)
    ws = wb.active

    headers: list[str] = []
    for col in range(2, (ws.max_column or 1) + 1):
        headers.append(_settlement_header(ws.cell(1, col).value))

    created_vehicles = 0
    created_settlements = 0
    updated_prices = 0

    for row in range(2, (ws.max_row or 1) + 1):
        vehicle_name = ws.cell(row, 1).value
        if vehicle_name is None or str(vehicle_name).strip() == "":
            continue
        vname = str(vehicle_name).strip()
        before = repo.get_vehicle_by_name(vname)
        vehicle = repo.ensure_vehicle(vname)
        if before is None:
            created_vehicles += 1

        for col, sett_name in enumerate(headers, start=2):
            if not sett_name:
                continue
            existing = repo.get_settlement_by_name(sett_name)
            settlement = repo.ensure_settlement(sett_name)
            if existing is None:
                created_settlements += 1

            cell_val = ws.cell(row, col).value
            if cell_val is None or str(cell_val).strip() == "":
                continue
            try:
                amount = float(str(cell_val).replace(",", ".").replace(" ", ""))
            except ValueError:
                continue
            repo.set_price(vehicle.id, settlement.id, amount)
            updated_prices += 1

    return {
        "created_vehicles": created_vehicles,
        "created_settlements": created_settlements,
        "updated_prices": updated_prices,
    }
