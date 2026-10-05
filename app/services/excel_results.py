from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image as PILImage

from app.db.models import Entry

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)


def export_results_xlsx(
    entries: list[Entry],
    output_path: Path,
    *,
    timezone_name: str = "Europe/Moscow",
) -> Path:
    tz = ZoneInfo(timezone_name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"

    headers = ["Дата", "Время", "Поселок", "Участок", "Тип авто", "Стоимость", "Фото"]
    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(1, idx, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 18

    for row_idx, entry in enumerate(entries, start=2):
        dt = entry.created_at
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo("UTC"))
        local = dt.astimezone(tz)
        sett = entry.settlement.name if entry.settlement else ""
        plot_name = entry.plot_name or ""
        vehicle_name = entry.vehicle.name if entry.vehicle else ""

        ws.cell(row_idx, 1, local.strftime("%d.%m.%Y"))
        ws.cell(row_idx, 2, local.strftime("%H:%M:%S"))
        ws.cell(row_idx, 3, sett)
        ws.cell(row_idx, 4, plot_name)
        ws.cell(row_idx, 5, vehicle_name)
        ws.cell(row_idx, 6, entry.price_amount)

        photo = Path(entry.photo_path)
        if photo.is_file():
            try:
                thumb = _make_thumb(photo, output_path.parent / f".thumb_{entry.id}.jpg")
                img = XLImage(str(thumb))
                img.width = 120
                img.height = 90
                ws.row_dimensions[row_idx].height = 72
                ws.add_image(img, f"G{row_idx}")
            except Exception:
                ws.cell(row_idx, 7, str(photo))
        else:
            ws.cell(row_idx, 7, "нет файла")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    for p in output_path.parent.glob(".thumb_*.jpg"):
        try:
            p.unlink()
        except OSError:
            pass
    return output_path


def _make_thumb(src: Path, dest: Path) -> Path:
    with PILImage.open(src) as im:
        im = im.convert("RGB")
        im.thumbnail((240, 180))
        dest.parent.mkdir(parents=True, exist_ok=True)
        im.save(dest, format="JPEG", quality=85)
    return dest


def parse_date_ddmmyyyy(value: str) -> datetime | None:
    value = (value or "").strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def parse_date_range(value: str) -> tuple[datetime, datetime] | None:
    raw = (value or "").strip().replace(" ", "")
    if "-" not in raw:
        return None
    left, right = raw.split("-", 1)
    d1 = parse_date_ddmmyyyy(left)
    d2 = parse_date_ddmmyyyy(right)
    if not d1 or not d2:
        return None
    return d1, d2
